"""Prepara tabelas relacionais, JSON, SQLite e Excel a partir dos snapshots.

Nenhum atributo de combate é inventado. Valores originais e problemas de
qualidade permanecem disponíveis para auditoria.
"""
import collections
import csv
import datetime as dt
import hashlib
import json
import pathlib
import re
import sqlite3
import sys
import unicodedata

from coletar_fontes import ROOT, RAW

OUT = ROOT / 'data' / 'processed'
CSV = OUT / 'csv'
JSON = OUT / 'json'
DOCS = ROOT / 'docs'
for folder in [OUT,CSV,JSON,DOCS]:
    folder.mkdir(parents=True,exist_ok=True)

TABLES = collections.defaultdict(list)
ISSUES = []
SOURCE = 'Dattebayo'
SOURCE_URL = 'https://github.com/viniciusschuelter/dattebayo-api/tree/a3871a6c931cd1f503af0bc16c80007c3a9ecda9/src/data'

def read(path):
    return json.loads(path.read_text(encoding='utf-8-sig'))

def norm(value):
    return re.sub(r'[^a-z0-9]','',unicodedata.normalize('NFKD',str(value)).encode('ascii','ignore').decode().lower())

def roman(value):
    value = unicodedata.normalize('NFKD',str(value)).encode('ascii','ignore').decode().lower()
    value = re.sub(r'[^a-z0-9 ]',' ',value)
    value = re.sub(r'ou|oo','o',value)
    value = value.replace('uu','u')
    return ' '.join(sorted(value.split()))

def values(value):
    if value is None:
        return []
    return value if isinstance(value,list) else [value]

def text(value):
    return ' | '.join(str(x) for x in values(value) if not isinstance(x,(dict,list))) or None

def uid(prefix,value):
    return prefix + '_' + hashlib.sha256(norm(value).encode()).hexdigest()[:16]

def issue(cid,field,reason,value=None):
    ISSUES.append({'personagem_id':cid,'campo':field,'motivo':reason,'valor_original':json.dumps(value,ensure_ascii=False) if value is not None else None})

def annotations(value):
    return ' | '.join(re.findall(r'\(([^()]*)\)',str(value))) or None

def label(value):
    return re.sub(r'\s+',' ',re.sub(r'\([^()]*\)','',str(value))).strip()

def flat(value,path=''):
    if isinstance(value,dict):
        for key,item in value.items():
            yield from flat(item,f'{path}.{key}' if path else key)
    elif isinstance(value,list):
        for i,item in enumerate(value):
            yield from flat(item,f'{path}[{i}]')
    else:
        yield path,value

def number_range(value):
    value = str(value)
    if re.search(r'unknown|varies|presumed|deceased|several',value,re.I):
        return None,None
    # Commas between three-digit groups are thousands separators in this source.
    value = re.sub(r'(?<=\d),(?=\d{3}(?:\D|$))','',value)
    nums = [float(n) for n in re.findall(r'\d+(?:\.\d+)?',value)]
    return (min(nums),max(nums)) if nums else (None,None)

def date_parts(value):
    try:
        date = dt.datetime.strptime(value,'%B %d')
        return date.month,date.day
    except (ValueError,TypeError):
        return None,None

def build():
    chars = read(RAW/'characters.json')
    char_by_id = {c['id']:c for c in chars}
    exact,loose = collections.defaultdict(set),collections.defaultdict(set)
    for char in chars:
        exact[norm(char['name'])].add(char['id'])
        loose[roman(char['name'])].add(char['id'])

    def resolve(name,aliases=()):
        for method,candidates in [('nome_exato_normalizado',exact.get(norm(name),set())),
                                  ('transliteracao_e_ordem',loose.get(roman(name),set()))]:
            if len(candidates)==1:
                return next(iter(candidates)),method
            if candidates:
                return None,'ambiguo'
        for index,method in [(exact,'alias_exato'),(loose,'alias_transliterado')]:
            keyer = norm if index is exact else roman
            candidates=set().union(*(index.get(keyer(alias),set()) for alias in aliases))
            if len(candidates)==1:
                return next(iter(candidates)),method
            if candidates:
                return None,'ambiguo'
        return None,'nao_resolvido'

    assets = read(ROOT/'assets'/'manifest.json') if (ROOT/'assets'/'manifest.json').exists() else []
    local_assets = {(x['kind'],x['source_id']):x for x in assets if x['status']=='downloaded'}
    officials={}
    official_native=collections.defaultdict(set)
    if (RAW/'narutop99'/'characters.json').exists():
        for item in read(RAW/'narutop99'/'characters.json'):
            oid=int(item['id'])
            name=item.get('name_en') or item['name']
            cid,method=resolve(name)
            if cid is None and method!='ambiguo':
                cid,method=resolve(label(name))
            officials[oid]={'raw':item,'personagem_id':cid,'method':method}
            if cid is not None:
                official_native[item['name']].add(cid)
    ani_chars,media_chars,voice_actors,casting = {},[],{},[]
    for path in sorted((RAW/'anilist').glob('characters_*.json')):
        media = read(path)['data']['Media']
        for edge in media['characters']['edges']:
            node = edge['node']
            ani_chars[node['id']] = node
            media_chars.append((media['id'],node['id'],edge['role']))
            for actor in edge.get('voiceActors',[]):
                voice_actors[actor['id']] = actor
                casting.append((media['id'],node['id'],actor['id'],actor.get('languageV2')))

    matched = {}
    for aid,node in sorted(ani_chars.items()):
        cid,method = resolve(node['name']['full'],node['name'].get('alternative',[]))
        if cid is None:
            native_candidates=official_native.get(node['name'].get('native'),set())
            if len(native_candidates)==1:
                cid=next(iter(native_candidates))
                method='nome_nativo_via_narutop99'
        matched[aid]=cid
        asset=local_assets.get(('characters',aid),{})
        birth=node.get('dateOfBirth') or {}
        TABLES['personagens_anilist'].append({'anilist_id':aid,'personagem_id':cid,'metodo_correspondencia':method,
            'nome':node['name']['full'],'nome_nativo':node['name'].get('native'),'genero':node.get('gender'),
            'idade_original':node.get('age'),'tipo_sanguineo':node.get('bloodType'),'nascimento_mes':birth.get('month'),
            'nascimento_dia':birth.get('day'),'favoritos_anilist':node.get('favourites'),'url_fonte':node.get('siteUrl'),
            'imagem_url':node.get('image',{}).get('large'),'imagem_local':asset.get('file'),'fonte':'AniList'})
        for spoiler,key in [(False,'alternative'),(True,'alternativeSpoiler')]:
            for alternative in node['name'].get(key,[]) or []:
                TABLES['aliases'].append({'anilist_id':aid,'personagem_id':cid,'alias':alternative,'spoiler':spoiler,'fonte':'AniList'})

    ani_by_char=collections.defaultdict(list)
    for aid,cid in matched.items():
        if cid is not None:
            ani_by_char[cid].append(aid)
    official_by_char=collections.defaultdict(list)
    for oid,item in officials.items():
        if item['personagem_id'] is not None:
            official_by_char[item['personagem_id']].append(oid)
    votes_global={int(r['character']):r for r in read(RAW/'narutop99'/'ranking_all.json')} if officials else {}
    for char in chars:
        cid=char['id']
        personal=char.get('personal',{})
        if not isinstance(personal,dict):
            issue(cid,'personal','estrutura_invalida',personal)
            personal={}
        rank=char.get('rank',{})
        if not isinstance(rank,dict):
            issue(cid,'rank','estrutura_invalida',rank)
            rank={}
        month,day=date_parts(personal.get('birthdate'))
        aids=ani_by_char.get(cid,[])
        best=max(aids,key=lambda x:ani_chars[x].get('favourites') or 0) if aids else None
        portrait=local_assets.get(('characters',best),{})
        oids=official_by_char.get(cid,[])
        official_best=min(oids,key=lambda x:votes_global.get(x,{}).get('rank',9999)) if oids else None
        official_image=local_assets.get(('official_portraits',official_best),{})
        if not portrait:
            portrait=official_image
        global_vote=votes_global.get(official_best,{})
        row={'personagem_id':cid,'nome':char['name'],'sexo_original':personal.get('sex'),
             'sexo_pt':{'Male':'Masculino','Female':'Feminino'}.get(personal.get('sex')),
             'especie_original':personal.get('species'),'status_original':personal.get('status'),
             'status_informado':'status' in personal,'aniversario_original':personal.get('birthdate'),
             'aniversario_mes':month,'aniversario_dia':day,'tipo_sanguineo':personal.get('bloodType'),
             'registro_ninja':rank.get('ninjaRegistration'),'cla_resumo':text(personal.get('clan')),
             'afiliacoes_resumo':text(personal.get('affiliation')),'equipes_resumo':text(personal.get('team')),
             'classificacao_resumo':text(personal.get('classification')),
             'quantidade_jutsus_listados_fonte':len(char.get('jutsu',[])),
             'quantidade_naturezas_listadas':len(char.get('natureType',[])),
             'quantidade_imagens_referenciadas':len(char.get('images',[])),
             'anilist_id_principal':best,'favoritos_anilist':ani_chars.get(best,{}).get('favourites'),
             'narutop99_id_principal':official_best,'posicao_narutop99':global_vote.get('rank'),
             'votos_narutop99':global_vote.get('score'),
             'imagem_url':portrait.get('url') or next(iter(char.get('images',[])),None),
             'imagem_local':portrait.get('file'),'fonte':SOURCE,'url_fonte':SOURCE_URL}
        TABLES['personagens'].append(row)
        for path,value in flat(char):
            TABLES['atributos_originais'].append({'personagem_id':cid,'caminho':path,
                'valor':str(value) if value is not None else None,'tipo_original':type(value).__name__,'fonte':SOURCE})
        known={'id','name','personal','images','debut','rank','family','jutsu','natureType','voiceActors','tools','uniqueTraits'}
        for key in char.keys()-known:
            issue(cid,key,'campo_inesperado_preservado_no_bruto',char[key])
        for attribute,value in personal.items():
            if attribute in ['age','height','weight'] and isinstance(value,dict):
                for context,original in value.items():
                    low,high=number_range(original)
                    TABLES['medidas_por_fase'].append({'personagem_id':cid,'atributo':{'age':'idade','height':'altura','weight':'peso'}[attribute],
                        'contexto_original':context,'valor_original':str(original),'valor_minimo':low,'valor_maximo':high,
                        'unidade':{'age':'anos','height':'cm','weight':'kg'}[attribute],
                        'limite_aberto':bool(re.search(r'\+|over|more|at least',str(original),re.I)),
                        'derivado_por_regex':True,'fonte':SOURCE})
            else:
                for pos,item in enumerate(values(value)):
                    TABLES['caracteristicas'].append({'personagem_id':cid,'caracteristica':attribute,'ordem':pos+1,
                        'valor_original':str(item),'fonte':SOURCE})
        for pos,url in enumerate(char.get('images',[]),1):
            TABLES['imagens_referenciadas'].append({'personagem_id':cid,'ordem':pos,'url':url,
                'fonte':SOURCE,'arquivo_local':None,'status':'referencia_nao_baixada'})
        for kind,field,prefix,dimension,bridge in [
            ('jutsu','jutsu','jutsu','jutsus','personagem_jutsu'),
            ('natureza','natureType','natureza','naturezas','personagem_natureza'),
            ('ferramenta','tools','ferramenta','ferramentas','personagem_ferramenta')]:
            for original in char.get(field,[]):
                if re.search(r'<[^>]+>',original) or (original.startswith('Naruto') and re.search(r'\d{4}',original)):
                    issue(cid,field,'conteudo_contaminado_separado_da_base_limpa',original)
                    TABLES['itens_quarentena'].append({'personagem_id':cid,'campo':field,'valor_original':original,
                        'motivo':'HTML ou catálogo de jogos no lugar de um atributo','fonte':SOURCE})
                    continue
                clean=label(original)
                item_id=uid(prefix,clean)
                TABLES[bridge].append({'personagem_id':cid,f'{kind}_id':item_id,'nome_original':original,
                    'observacoes':annotations(original),'fonte':SOURCE})
                TABLES[dimension].append({f'{kind}_id':item_id,'nome':clean,'fonte':SOURCE})
        row['quantidade_jutsus_utilizaveis']=len({r['jutsu_id'] for r in TABLES['personagem_jutsu'] if r['personagem_id']==cid})
        for feature in char.get('uniqueTraits',[]):
            TABLES['tracos_especiais'].append({'personagem_id':cid,'traco_original':feature,'fonte':SOURCE})
        for relationship,target in char.get('family',{}).items():
            resolved,method=resolve(target)
            TABLES['familia'].append({'personagem_id':cid,'parentesco_original':relationship,'nome_relacionado_original':target,
                'personagem_relacionado_id':resolved,'metodo_correspondencia':method,'fonte':SOURCE})
        for medium,reference in char.get('debut',{}).items():
            if medium=='appearsIn':
                for item in reference.split(','):
                    TABLES['aparicoes_por_midia'].append({'personagem_id':cid,'midia':item.strip(),'fonte':SOURCE})
            else:
                TABLES['estreias'].append({'personagem_id':cid,'midia':medium,'referencia_original':reference,'fonte':SOURCE})
        ninja_rank=rank.get('ninjaRank',{})
        if isinstance(ninja_rank,dict):
            for context,value in ninja_rank.items():
                TABLES['patentes_por_fase'].append({'personagem_id':cid,'contexto_original':context,'patente_original':value,'fonte':SOURCE})
        for language,actors in char.get('voiceActors',{}).items():
            for actor in values(actors):
                TABLES['dublagem_original'].append({'personagem_id':cid,'idioma_original':language,'nome_credito_original':actor,'fonte':SOURCE})

    for dimension,idcol,filename in [('clas','cla_id','clans'),('vilas','vila_id','villages'),
                                     ('equipes','equipe_id','teams'),('kekkei_genkai','kekkei_genkai_id','kekkei-genkai')]:
        for item in read(RAW/(filename+'.json')):
            TABLES[dimension].append({idcol:item['id'],'nome':item['name'],'fonte':SOURCE})
            for cid in item.get('characters',[]):
                if cid not in char_by_id:
                    issue(cid,filename,'id_de_personagem_nao_encontrado',item['id'])
                TABLES['personagem_'+dimension].append({'personagem_id':cid,idcol:item['id'],'fonte':SOURCE})
    for kind,file in [('bijuus','tailed-beasts'),('akatsuki','akatsuki'),('kara','kara')]:
        for item in read(RAW/(file+'.json')):
            # Estas coleções incluem associados, variantes e criaturas; não provam filiação canônica.
            TABLES['colecoes_fonte'].append({'colecao':kind,'personagem_id':item['id'],'nome_na_colecao':item['name'],
                'encontrado_na_base_principal':item['id'] in char_by_id,'fonte':SOURCE})
    symbol_manifest=ROOT/'assets'/'symbols'/'manifest.json'
    if symbol_manifest.exists():
        slug_village={'konoha':'Konohagakure','suna':'Sunagakure','kiri':'Kirigakure','kumo':'Kumogakure','iwa':'Iwagakure'}
        village_ids={r['nome']:r['vila_id'] for r in TABLES['vilas']}
        for item in read(symbol_manifest):
            TABLES['simbolos_vilas'].append({'vila_id':village_ids.get(slug_village[item['slug']]),
                'nome_vila':slug_village[item['slug']],'arquivo_local':item['file'],
                'url_fonte':item['url'],'sha256':item['sha256'],'fonte':'NARUTOP99'})

    for mid,aid,role in sorted(set(media_chars)):
        TABLES['elenco_por_serie'].append({'serie_anilist_id':mid,'anilist_id':aid,'personagem_id':matched[aid],
            'papel_original':role,'fonte':'AniList'})
    for aid,actor in sorted(voice_actors.items()):
        TABLES['dubladores'].append({'dublador_anilist_id':aid,'nome':actor['name']['full'],
            'nome_nativo':actor['name'].get('native'),'idioma':actor.get('languageV2'),
            'imagem_url':actor.get('image',{}).get('large'),'fonte':'AniList'})
    for mid,aid,vid,language in sorted(set(casting),key=str):
        TABLES['creditos_dublagem'].append({'serie_anilist_id':mid,'anilist_id':aid,'personagem_id':matched[aid],
            'dublador_anilist_id':vid,'idioma':language,'fonte':'AniList'})

    media_all={}
    for path in sorted((RAW/'anilist').glob('search*.json')):
        for item in read(path)['data']['Page']['media']:
            media_all[item['id']]=item
    for path in sorted((RAW/'anilist').glob('media_*.json')):
        item=read(path)['data']['Media']
        media_all[item['id']]=item
    for mid,item in sorted(media_all.items()):
        cover=local_assets.get(('covers',mid),{})
        banner=local_assets.get(('banners',mid),{})
        TABLES['series_filmes'].append({'serie_anilist_id':mid,'mal_id':item.get('idMal'),
            'titulo_romaji':item['title'].get('romaji'),'titulo_ingles':item['title'].get('english'),
            'titulo_japones':item['title'].get('native'),'formato':item.get('format'),'status':item.get('status'),
            'episodios':item.get('episodes'),'duracao_episodio_minutos':item.get('duration'),
            'ano_temporada':item.get('seasonYear'),'nota_media_0_100':item.get('averageScore'),
            'usuarios_anilist':item.get('popularity'),'favoritos_anilist':item.get('favourites'),
            'inicio_ano':(item.get('startDate') or {}).get('year'),'inicio_mes':(item.get('startDate') or {}).get('month'),
            'inicio_dia':(item.get('startDate') or {}).get('day'),'fim_ano':(item.get('endDate') or {}).get('year'),
            'fim_mes':(item.get('endDate') or {}).get('month'),'fim_dia':(item.get('endDate') or {}).get('day'),
            'capa_local':cover.get('file'),'banner_local':banner.get('file'),'capa_url':cover.get('url'),
            'banner_url':item.get('bannerImage'),'url_fonte':item.get('siteUrl'),'fonte':'AniList'})
        for genre in item.get('genres',[]):
            TABLES['generos_por_serie'].append({'serie_anilist_id':mid,'genero':genre,'fonte':'AniList'})
        for tag in item.get('tags',[]):
            TABLES['tags_por_serie'].append({'serie_anilist_id':mid,'tag_id':tag['id'],'nome':tag['name'],
                'categoria':tag.get('category'),'relevancia_0_100':tag.get('rank'),
                'spoiler':bool(tag.get('isMediaSpoiler') or tag.get('isGeneralSpoiler')),'fonte':'AniList'})
        for relation in item.get('relations',{}).get('edges',[]):
            target=relation['node']
            TABLES['relacoes_entre_obras'].append({'serie_anilist_id':mid,'obra_relacionada_id':target['id'],
                'tipo_relacao':relation['relationType'],'tipo_obra':target.get('type'),
                'titulo':target['title']['romaji'],'url_fonte':target.get('siteUrl'),'fonte':'AniList'})
        for score in item.get('stats',{}).get('scoreDistribution',[]):
            TABLES['distribuicao_notas'].append({'serie_anilist_id':mid,'nota':score['score'],'usuarios':score['amount'],'fonte':'AniList'})
        for status in item.get('stats',{}).get('statusDistribution',[]):
            TABLES['distribuicao_status'].append({'serie_anilist_id':mid,'status_lista':status['status'],'usuarios':status['amount'],'fonte':'AniList'})
        for link in item.get('externalLinks',[]):
            TABLES['links_por_serie'].append({'serie_anilist_id':mid,'site':link['site'],'url':link['url'],
                'tipo':link.get('type'),'idioma':link.get('language'),'fonte':'AniList'})

    serie_ids={'naruto':20,'naruto-shippuden':1735,'boruto':97938}
    for serie in read(RAW/'episode_source'/'anime.json'):
        slug=serie['slug']
        if slug not in serie_ids:
            continue
        arcs=collections.defaultdict(list)
        prev_date=None
        for ep in serie['episodes']:
            arc_id=uid('arco',slug+' '+str(ep.get('arc','Sem arco')))
            air_date=ep.get('airDate')
            quality='nao_verificada'
            if air_date and prev_date and air_date<prev_date:
                quality='inconsistencia_cronologica'
                issue(None,'episodio.data_exibicao','data_anterior_ao_episodio_precedente',{'serie':slug,'episodio':ep['number'],'data':air_date})
            if air_date:
                prev_date=air_date
            TABLES['episodios'].append({'episodio_id':slug+'_'+str(ep['number']).zfill(3),'serie_anilist_id':serie_ids[slug],
                'serie':slug,'numero':ep['number'],'titulo_original':ep.get('title'),'classificacao_comunitaria':ep.get('type'),
                'arco_id':arc_id,'data_exibicao_original':air_date,'data_exibicao_validada':None,
                'qualidade_data':quality,'fonte':'SkipTheFiller','url_fonte':serie.get('url'),
                'licenca':'CC BY 4.0'})
            arcs[ep.get('arc','Sem arco')].append(ep['number'])
        for arc,numbers in arcs.items():
            TABLES['arcos'].append({'arco_id':uid('arco',slug+' '+arc),'serie_anilist_id':serie_ids[slug],
                'nome':arc,'primeiro_episodio':min(numbers),'ultimo_episodio':max(numbers),
                'quantidade_episodios':len(numbers),'fonte':'SkipTheFiller','licenca':'CC BY 4.0'})
    community=read(RAW/'episode_source'/'characters.json')
    for slug,mid in serie_ids.items():
        for char in community.get(slug,[]):
            cid,method=resolve(char['name'])
            TABLES['papeis_comunitarios'].append({'serie_anilist_id':mid,'personagem_id':cid,'nome_original':char['name'],
                'papel_original':char.get('role'),'primeiro_episodio_informado':char.get('firstEp'),
                'metodo_correspondencia':method,'url_fonte':char.get('url'),'fonte':'SkipTheFiller','licenca':'CC BY 4.0'})

    for item in assets:
        TABLES['arquivos_imagem'].append({'tipo':item['kind'],'fonte_id':item['source_id'],
            'personagem_id':matched.get(item['source_id']) if item['kind']=='characters' else officials.get(item['source_id'],{}).get('personagem_id') if item['kind'] in ['official_portraits','official_cards'] else None,
            'nome':item['name'],'url':item['url'],'arquivo_local':item.get('file'),'status':item['status'],
            'largura':item.get('width'),'altura':item.get('height'),'formato':item.get('format'),
            'canal_alpha':item.get('has_alpha_channel'),'bytes':item.get('bytes'),'sha256':item.get('sha256'),
            'fonte':item['source'],'verificado_em_utc':item['checked_at_utc']})

    villages={'木ノ葉':'Konohagakure','砂':'Sunagakure','霧':'Kirigakure','雲':'Kumogakure','岩':'Iwagakure','その他':'Outros'}
    for oid,item in officials.items():
        raw=item['raw']
        image=local_assets.get(('official_portraits',oid),{})
        card=local_assets.get(('official_cards',oid),{})
        TABLES['personagens_narutop99'].append({'narutop99_id':oid,'personagem_id':item['personagem_id'],
            'metodo_correspondencia':item['method'],'nome':raw.get('name_en'),'nome_japones':raw['name'],
            'vila_japones':raw.get('village'),'vila':villages.get(raw.get('village'),raw.get('village')),
            'retrato_local':image.get('file'),'card_local':card.get('file'),
            'url_fonte':'https://narutop99.naruto-official.com/en/','fonte':'NARUTOP99'})
    for path in sorted((RAW/'narutop99').glob('ranking_*.json')):
        region=path.stem.removeprefix('ranking_')
        for vote in read(path):
            oid=int(vote['character'])
            TABLES['votacao_narutop99'].append({'narutop99_id':oid,'personagem_id':officials[oid]['personagem_id'],
                'regiao':region,'posicao':vote['rank'],'votos':vote['score'],
                'periodo_inicio':'2022-12-17','periodo_fim':'2023-01-31','resultado_ano':2023,
                'fonte':'NARUTOP99','url_fonte':'https://narutop99.naruto-official.com/en/'})
    official_html=RAW/'narutop99'/'index.html'
    if official_html.exists():
        from html_dados import demographics
        for row in demographics(official_html.read_text(encoding='utf-8')):
            row['personagem_id']=officials[row['narutop99_id']]['personagem_id']
            row['fonte']='NARUTOP99'
            row['resultado_ano']=2023
            TABLES['ranking_demografico'].append(row)

    for dim,key in [('jutsus','jutsu_id'),('naturezas','natureza_id'),('ferramentas','ferramenta_id')]:
        unique={}
        for row in TABLES[dim]:
            unique.setdefault(row[key],row)
        TABLES[dim]=sorted(unique.values(),key=lambda r:r['nome'])
    for table in list(TABLES):
        seen=set(); unique=[]
        for row in TABLES[table]:
            signature=json.dumps(row,sort_keys=True,ensure_ascii=False)
            if signature not in seen:
                unique.append(row);seen.add(signature)
        TABLES[table]=unique
    TABLES['problemas_qualidade']=ISSUES
    return chars

def exports():
    schema={}
    db=sqlite3.connect(OUT/'naruto.sqlite')
    for table,rows in sorted(TABLES.items()):
        if not rows:
            continue
        fields=list(dict.fromkeys(k for row in rows for k in row))
        types={}
        for field in fields:
            typeset={type(r[field]) for r in rows if r.get(field) is not None}
            types[field]='boolean' if typeset and typeset<={bool} else ('integer' if typeset and typeset<={int} else ('number' if typeset and typeset<={int,float} else 'text'))
        schema[table]={'rows':len(rows),'columns':types}
        with (CSV/(table+'.csv')).open('w',encoding='utf-8-sig',newline='') as stream:
            writer=csv.DictWriter(stream,fieldnames=fields,delimiter=';')
            writer.writeheader()
            writer.writerows({k:('true' if v is True else 'false' if v is False else v) for k,v in row.items()} for row in rows)
        (JSON/(table+'.json')).write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
        db.execute(f'DROP TABLE IF EXISTS "{table}"')
        sqlite_types={'integer':'INTEGER','boolean':'INTEGER','number':'REAL','text':'TEXT'}
        columns=','.join(f'"{f}" {sqlite_types[types[f]]}' for f in fields)
        db.execute(f'CREATE TABLE "{table}" ({columns})')
        placeholders=','.join('?' for _ in fields)
        db.executemany(f'INSERT INTO "{table}" VALUES ({placeholders})',[[r.get(f) for f in fields] for r in rows])
    db.commit();db.close()
    (OUT/'schema.json').write_text(json.dumps(schema,ensure_ascii=False,indent=2),encoding='utf-8')
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill
        from openpyxl.cell import WriteOnlyCell
        workbook=Workbook(write_only=True)
        for table,meta in schema.items():
            sheet=workbook.create_sheet(table[:31])
            sheet.freeze_panes='A2'
            fields=list(meta['columns'])
            cells=[]
            for field in fields:
                cell=WriteOnlyCell(sheet,field)
                cell.font=Font(bold=True,color='FFFFFF')
                cell.fill=PatternFill('solid',fgColor='D96B16')
                cells.append(cell)
            sheet.append(cells)
            for row in TABLES[table]:
                cells=[]
                for field in fields:
                    value=row.get(field)
                    cell=WriteOnlyCell(sheet,value)
                    if isinstance(value,str):
                        cell.data_type='s'
                    cells.append(cell)
                sheet.append(cells)
            sheet.auto_filter.ref=f'A1:{__import__("openpyxl").utils.get_column_letter(len(fields))}{len(TABLES[table])+1}'
        workbook.save(OUT/'Naruto_BI.xlsx')
    except ImportError:
        print('openpyxl unavailable: CSV/JSON/SQLite exported; Excel skipped',flush=True)
    return schema

def report(chars,schema):
    asset_counts=collections.Counter(r['tipo'] for r in TABLES['arquivos_imagem'] if r['status']=='downloaded')
    coverage=[]
    for name,check in [
        ('Imagem referenciada',lambda c:bool(c.get('images'))),
        ('Retrato local associado',lambda c:any(r['personagem_id']==c['id'] and r.get('imagem_local') for r in TABLES['personagens'])),
        ('Jutsus listados',lambda c:bool(c.get('jutsu'))),('Naturezas de chakra',lambda c:bool(c.get('natureType'))),
        ('Clã',lambda c:isinstance(c.get('personal'),dict) and bool(c['personal'].get('clan'))),
        ('Altura por fase',lambda c:isinstance(c.get('personal'),dict) and bool(c['personal'].get('height'))),
        ('Família',lambda c:bool(c.get('family'))),('Patente',lambda c:isinstance(c.get('rank'),dict) and bool(c['rank'].get('ninjaRank'))),
        ('Estreia',lambda c:bool(c.get('debut')))]:
        count=sum(bool(check(c)) for c in chars)
        coverage.append({'campo':name,'personagens':count,'percentual':round(100*count/len(chars),2)})
    summary={'prepared_at_utc':dt.datetime.now(dt.timezone.utc).isoformat(),
        'personagens':len(chars),'tabelas':len(schema),'linhas_total':sum(x['rows'] for x in schema.values()),
        'anilist_personagens':len(TABLES['personagens_anilist']),
        'anilist_associados':sum(r['personagem_id'] is not None for r in TABLES['personagens_anilist']),
        'arquivos_imagem_baixados':dict(asset_counts),'imagens_referenciadas':len(TABLES['imagens_referenciadas']),
        'episodios':len(TABLES['episodios']),'jutsus_distintos':len(TABLES['jutsus']),
        'personagens_narutop99':len(TABLES['personagens_narutop99']),
        'votos_oficiais_globais':sum(r['votos'] for r in TABLES['votacao_narutop99'] if r['regiao']=='all'),
        'problemas_qualidade':len(ISSUES),'cobertura':coverage,'tabelas_linhas':{k:v['rows'] for k,v in schema.items()}}
    (OUT/'resumo.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    lines=['# Inventário da coleta Naruto','',
        'Coleta preparada em '+summary['prepared_at_utc']+'. Valores de popularidade são um retrato da coleta.',
        '',f'**{len(chars):,} personagens; {len(schema)} tabelas; {sum(asset_counts.values())} imagens locais verificadas.**','',
        '## Cobertura real dos personagens','', '| Campo | Personagens | Cobertura |','|---|---:|---:|']
    lines += [f'| {x["campo"]} | {x["personagens"]} | {x["percentual"]}% |' for x in coverage]
    lines += ['', '## Tabelas exportadas','','| Tabela | Linhas |','|---|---:|']
    lines += [f'| {k} | {v["rows"]} |' for k,v in schema.items()]
    lines += ['', '## Limites que afetam o BI','',
        '- Fontes comunitárias, com cobertura desigual e erros de extração. Ausência de status não significa que o personagem esteja vivo; ausência de espécie não prova que seja humano.',
        '- Os nomes dos jutsus estão disponíveis, mas não há descrições completas, custo de chakra ou escala de força validada. Contagem de técnicas não é medida de poder.',
        '- As coleções Akatsuki/Kara/bijūs da fonte incluem associados e variantes. Não tratar todas as entradas como membros oficiais ou como as nove bestas tradicionais.',
        '- Há nomes repetidos (ex.: Soul Tao); use personagem_id, nunca o nome como chave.',
        '- Dados de idade, altura, peso e patentes pertencem a contextos distintos. Não combinar fases para apresentar uma ficha atual.',
        '- Foram encontrados 11 campos personal e 2 campos rank com estrutura inválida. Estão preservados no bruto e sinalizados, sem inventar valores.',
        '- A coluna data_exibicao_original é uma transcrição não validada. A fonte tem erros cronológicos, inclusive o episódio 220 de Naruto com ano 2005. data_exibicao_validada fica vazia até confirmação por outra fonte.',
        '- Canon/Filler/Mixed e agrupamento em arcos são classificações editoriais da comunidade, não uma declaração oficial do estúdio.',
        '- As imagens originais do Fandom são referências: amostras retornaram HTTP 403. Não houve tentativa de contornar o bloqueio. As imagens efetivamente locais vieram da AniList e do NARUTOP99 oficial.',
        '- Correspondências entre fontes são documentadas por método. Casos ambíguos ficam sem personagem_id e podem ser vistos em personagens_anilist; nenhuma correspondência aproximada por similaridade foi aplicada.',
        '- Não foi verificada a classificação cânone/filler de cada personagem nem o estado atual dos personagens no mangá mais recente.',
        '- O NARUTOP99 é uma votação histórica de 2022/2023, não uma medida atual de popularidade nem de força. A região all contém os mesmos votos das oito regiões somadas: filtre all OU as regiões; nunca some ambos.',
        '- O ranking por idade/gênero contém somente os cinco primeiros de cada recorte. As porcentagens e os votos por recorte não foram inferidos.',
        '', '## Fontes','',
        '- [Dattebayo API — repositório e snapshot](https://github.com/viniciusschuelter/dattebayo-api/tree/a3871a6c931cd1f503af0bc16c80007c3a9ecda9): personagens e relações. Commit, hash e horário em data/raw/manifest.json.',
        '- [AniList API](https://docs.anilist.co/): elenco, retratos, créditos, séries e metadados. Snapshots e hashes em data/raw/anilist/manifest.json.',
        '- [NARUTOP99 oficial](https://narutop99.naruto-official.com/en/): catálogo de 488 participantes, votos globais e em oito regiões, top 5 por idade/gênero e imagens originais do site. Arrays estáticos públicos foram interpretados como JSON sem execução de JavaScript. Arquivos e hashes em data/raw/narutop99/manifest.json.',
        '- [SkipTheFiller](https://skipthefiller.com), via [dataset aberto](https://github.com/animecannonhunter-hash/anime-filler-list-data/tree/a7b28ff590b70652843125b77b9a80f334d2c3ca): episódios, arcos e papéis comunitários. [CC BY 4.0](https://creativecommons.org/licenses/by/4.0/). Foram filtradas as três séries e normalizados nomes de campos; os valores originais permanecem nos snapshots.',
        '', 'As imagens mantêm os direitos dos respectivos autores/titulares; a licença do código de uma API não concede automaticamente direitos sobre a arte da franquia.',
        '', '## Material adicional pesquisado','',
        '- NarutoDB: domínio respondeu HTML e não foi usado como API funcional.',
        '- Jikan: conexões expiraram; não é uma fonte desta entrega.',
        '- data/raw/naruto_secondary.json: lista adicional de referências de imagens de muhammadpauzi/naruto-api, commit 18b04502ee1e9606e79f6b190b2ced15e8f6849f; não incorporada às tabelas para evitar identidades ambíguas.',
        '- Não foram incorporados PDFs não verificados de databooks nem notas de força criadas por fãs.']
    (DOCS/'INVENTARIO.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    print(json.dumps({k:v for k,v in summary.items() if k not in ['cobertura','tabelas_linhas']},ensure_ascii=True,indent=2),flush=True)

if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    chars=build()
    schema=exports()
    report(chars,schema)
