"""Auditoria da entrega: integridade das tabelas, imagens e votação oficial."""
import collections
import csv
import hashlib
import json
import pathlib
import sqlite3
import sys
from coletar_fontes import ROOT

DATA=ROOT/'data'/'processed'
def read(path):
    return json.loads(path.read_text(encoding='utf-8-sig'))

def relationships():
    rows=[]
    def add(parent,key,child,foreign=None):
        rows.append({'tabela_um':parent,'coluna_um':key,'tabela_muitos':child,
                     'coluna_muitos':foreign or key,'cardinalidade':'1:N','direcao':'unica'})
    for child in ['personagem_clas','personagem_vilas','personagem_equipes','personagem_kekkei_genkai',
                  'personagem_jutsu','personagem_natureza','personagem_ferramenta','familia','medidas_por_fase',
                  'patentes_por_fase','caracteristicas','estreias','aparicoes_por_midia','dublagem_original',
                  'tracos_especiais','imagens_referenciadas','personagens_anilist','personagens_narutop99',
                  'papeis_comunitarios','arquivos_imagem']:
        add('personagens','personagem_id',child)
    for dim,key,child in [('clas','cla_id','personagem_clas'),('vilas','vila_id','personagem_vilas'),
        ('equipes','equipe_id','personagem_equipes'),('kekkei_genkai','kekkei_genkai_id','personagem_kekkei_genkai'),
        ('jutsus','jutsu_id','personagem_jutsu'),('naturezas','natureza_id','personagem_natureza'),
        ('ferramentas','ferramenta_id','personagem_ferramenta'),('vilas','vila_id','simbolos_vilas'),
        ('personagens_anilist','anilist_id','aliases'),('personagens_anilist','anilist_id','elenco_por_serie'),
        ('personagens_anilist','anilist_id','creditos_dublagem'),('dubladores','dublador_anilist_id','creditos_dublagem'),
        ('series_filmes','serie_anilist_id','elenco_por_serie'),('series_filmes','serie_anilist_id','creditos_dublagem'),
        ('series_filmes','serie_anilist_id','arcos'),('arcos','arco_id','episodios'),
        ('personagens_narutop99','narutop99_id','votacao_narutop99'),
        ('personagens_narutop99','narutop99_id','ranking_demografico')]:
        add(dim,key,child)
    return rows

def main():
    schema=read(DATA/'schema.json')
    tables={name:read(DATA/'json'/(name+'.json')) for name in schema}
    checks=[]
    for name,meta in schema.items():
        rows=tables[name]
        assert len(rows)==meta['rows'],f'JSON row count: {name}'
        with (DATA/'csv'/(name+'.csv')).open(encoding='utf-8-sig',newline='') as stream:
            reader=csv.DictReader(stream,delimiter=';')
            assert reader.fieldnames==list(meta['columns']),f'CSV columns: {name}'
            assert sum(1 for _ in reader)==len(rows),f'CSV rows: {name}'
    checks.append(f'{len(schema)} tabelas JSON/CSV com contagens e cabeçalhos consistentes')
    database=sqlite3.connect(DATA/'naruto.sqlite')
    assert database.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
    for name,rows in tables.items():
        assert database.execute(f'SELECT COUNT(*) FROM "{name}"').fetchone()[0]==len(rows)
    database.close()
    checks.append('SQLite íntegro, com as mesmas contagens de linhas')
    rels=relationships()
    for rel in rels:
        parent=tables[rel['tabela_um']];child=tables[rel['tabela_muitos']]
        ids=[r[rel['coluna_um']] for r in parent]
        assert None not in ids and len(ids)==len(set(ids)),f'Unique key: {rel}'
        missing={r[rel['coluna_muitos']] for r in child if r.get(rel['coluna_muitos']) is not None}-set(ids)
        assert not missing,f'Foreign keys {rel}: {missing}'
    with (ROOT/'powerbi'/'relacionamentos.csv').open('w',encoding='utf-8-sig',newline='') as stream:
        writer=csv.DictWriter(stream,fieldnames=list(rels[0]),delimiter=';');writer.writeheader();writer.writerows(rels)
    checks.append(f'{len(rels)} relações propostas com chaves únicas e sem referências órfãs não nulas')
    charids={r['personagem_id'] for r in tables['personagens']}
    assert len(charids)==1431
    assert all(r['personagem_relacionado_id'] is None or r['personagem_relacionado_id'] in charids for r in tables['familia'])
    assert all(not r['nome'].lstrip().startswith('<') and '2016' not in r['nome'] for r in tables['jutsus'])
    checks.append('Personagens e parentesco consistentes; resíduos HTML/catálogos separados dos jutsus')
    expected={20:220,1735:500,97938:293}
    actual=collections.Counter(r['serie_anilist_id'] for r in tables['episodios'])
    assert actual==expected
    assert len({r['episodio_id'] for r in tables['episodios']})==1013
    for mid,total in expected.items():
        assert {r['numero'] for r in tables['episodios'] if r['serie_anilist_id']==mid}==set(range(1,total+1))
    checks.append('1.013 episódios, sem números faltantes ou duplicados nas três séries')
    global_votes={r['narutop99_id']:r['votos'] for r in tables['votacao_narutop99'] if r['regiao']=='all'}
    regional_votes=collections.Counter()
    for row in tables['votacao_narutop99']:
        if row['regiao']!='all':
            regional_votes[row['narutop99_id']]+=row['votos']
    assert regional_votes==global_votes
    assert len(global_votes)==488 and sum(global_votes.values())==4598612
    assert len(tables['ranking_demografico'])==120
    checks.append('4.598.612 votos oficiais: a soma das oito regiões confere com o global para cada um dos 488 participantes')
    assets=read(ROOT/'assets'/'manifest.json')
    for asset in assets:
        assert asset['status']=='downloaded',asset
        path=ROOT/asset['file']
        assert path.is_relative_to(ROOT) and path.exists()
        assert hashlib.sha256(path.read_bytes()).hexdigest()==asset['sha256']
        assert asset['width']>0 and asset['height']>0
    checks.append(f'{len(assets)} imagens locais com dimensões, existência e SHA-256 conferidos')
    for row in tables['personagens']:
        if row.get('imagem_local'):
            assert (ROOT/row['imagem_local']).exists()
        if row.get('registro_ninja'):
            assert isinstance(row['registro_ninja'],str)
    from openpyxl import load_workbook
    workbook=load_workbook(DATA/'Naruto_BI.xlsx',read_only=True,data_only=False)
    assert len(workbook.sheetnames)==len(schema)
    for name,meta in schema.items():
        sheet=workbook[name[:31]]
        iterator=sheet.iter_rows(values_only=True)
        assert next(iterator)==tuple(meta['columns']),f'Excel headers: {name}'
        assert sum(1 for _ in iterator)==meta['rows'],f'Excel rows: {name}'
    workbook.close()
    checks.append('Excel reaberto: todas as abas, cabeçalhos e contagens de linhas conferidos')
    result={'status':'passed','checks':checks,'limitations':['Sem execução de DAX ou Power Query no Power BI nesta etapa.','Datas de episódios e atributos de combate não foram validados como fatos oficiais.']}
    (DATA/'validacao.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False,indent=2),flush=True)

if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
